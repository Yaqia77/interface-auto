"""凭证变量（secret_vars.xlsx「全局变量」表）读写：值永不明文回显，仅掩码。

安全约定：
- GET 只返回掩码值 + 长度，绝不把真实凭证透给前端；
- 更新时 value 为空表示保持原值不变（前端无从回显原值，也就无从泄露）；
- 写回沿用「临时文件校验后原子替换」，文件被占用转 409。
"""
import os
from pathlib import Path

from openpyxl import Workbook, load_workbook

from core import excel_loader
from core.excel_loader import ExcelFormatError

SHEET = "全局变量"
COLS = ["变量名", "值", "备注"]

# 名称含这些片段视为敏感凭证：掩码时不暴露任何真实字符，仅给长度
_SENSITIVE_HINTS = ("password", "pwd", "token", "secret", "authorization", "cookie", "key", "auth")
# 名称含这些片段归类为“登录态/凭证”（可关联自动刷新）
_LOGIN_HINTS = ("token", "authorization", "cookie", "auth")
# 名称含这些片段归类为“密文”
_CIPHER_HINTS = ("password", "pwd", "secret", "key")


class SecretError(Exception):
    """kind: invalid（值非法）/ not_found（文件或变量不存在）/ conflict（重名）/ locked（被占用）"""

    def __init__(self, kind: str, message: str):
        super().__init__(message)
        self.kind = kind


def _is_sensitive(name: str) -> bool:
    n = name.lower()
    return any(h in n for h in _SENSITIVE_HINTS)


def _kind(name: str) -> str:
    n = name.lower()
    if any(h in n for h in _LOGIN_HINTS):
        return "登录态/凭证"
    if any(h in n for h in _CIPHER_HINTS):
        return "密文"
    return "固定值"


def _mask(name: str, value) -> tuple[str, int]:
    """返回 (掩码串, 原值长度)。敏感名一律全遮盖，仅露长度。"""
    s = "" if value is None else str(value)
    length = len(s)
    if length == 0:
        return "", 0
    if _is_sensitive(name):
        return "•" * 8 + f"（{length} 位）", length
    if length <= 6:
        return s[:1] + "****", length
    return s[:4] + "****" + s[-2:] + f"（{length} 位）", length


def read_all(path: Path) -> dict:
    """读取「全局变量」表并逐行掩码。文件不存在返回 {exists:false, items:[]}。"""
    if not path.exists():
        return {"exists": False, "items": []}
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        if SHEET not in wb.sheetnames:
            raise SecretError("invalid", f"凭证文件缺少工作表「{SHEET}」")
        rows = list(wb[SHEET].iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return {"exists": True, "items": []}

    idx = {"name": 0, "value": 1, "remark": 2}
    for i, cell in enumerate(rows[0]):
        h = str(cell).strip() if cell is not None else ""
        if h == "变量名":
            idx["name"] = i
        elif h == "值":
            idx["value"] = i
        elif h == "备注":
            idx["remark"] = i

    items = []
    for row in rows[1:]:
        raw_name = row[idx["name"]] if idx["name"] < len(row) else None
        name = str(raw_name).strip() if raw_name is not None else ""
        if not name:
            continue
        value = row[idx["value"]] if idx["value"] < len(row) else None
        remark = row[idx["remark"]] if idx["remark"] < len(row) else None
        masked, length = _mask(name, value)
        items.append({
            "name": name,
            "masked_value": masked,
            "length": length,
            "empty": length == 0,
            "kind": _kind(name),
            "remark": str(remark).strip() if remark is not None else "",
        })
    return {"exists": True, "items": items}


def _ensure(path: Path):
    """取可写工作簿 + 「全局变量」表 + 列定位；文件不存在则按表头新建。"""
    if path.exists():
        wb = load_workbook(path)
        if SHEET not in wb.sheetnames:
            raise SecretError("invalid", f"凭证文件缺少工作表「{SHEET}」")
        ws = wb[SHEET]
    else:
        path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET
        for c, h in enumerate(COLS, start=1):
            ws.cell(row=1, column=c, value=h)

    cols = {"name": 1, "value": 2, "remark": 3}
    for cell in ws[1]:
        h = str(cell.value).strip() if cell.value is not None else ""
        if h == "变量名":
            cols["name"] = cell.column
        elif h == "值":
            cols["value"] = cell.column
        elif h == "备注":
            cols["remark"] = cell.column
    return wb, ws, cols


def _find_row(ws, name_col: int, name: str) -> int | None:
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=name_col).value
        if v is not None and str(v).strip() == name:
            return r
    return None


def _existing_names(ws, name_col: int) -> set[str]:
    out = set()
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=name_col).value
        if v is not None and str(v).strip():
            out.add(str(v).strip())
    return out


def create_secret(path: Path, name: str, value: str, remark: str | None) -> None:
    name = (name or "").strip()
    if not name:
        raise SecretError("invalid", "变量名不能为空")
    if any(c in name for c in "\r\n"):
        raise SecretError("invalid", "变量名不能包含换行")
    if value is None:
        raise SecretError("invalid", "凭证值不能为空")
    wb, ws, cols = _ensure(path)
    try:
        if name in _existing_names(ws, cols["name"]):
            raise SecretError("conflict", f"变量 {name} 已存在，请直接编辑其值")
        r = ws.max_row + 1 if ws.max_row >= 1 else 2
        ws.cell(row=r, column=cols["name"]).value = name
        ws.cell(row=r, column=cols["value"]).value = value
        ws.cell(row=r, column=cols["remark"]).value = remark or ""
        _save(wb, path)
    finally:
        wb.close()


def update_secret(path: Path, name: str, value: str | None, remark: str | None) -> None:
    if not path.exists():
        raise SecretError("not_found", "凭证文件不存在")
    wb, ws, cols = _ensure(path)
    try:
        r = _find_row(ws, cols["name"], (name or "").strip())
        if not r:
            raise SecretError("not_found", f"变量不存在: {name}")
        if value is not None and str(value) != "":
            ws.cell(row=r, column=cols["value"]).value = value
        if remark is not None:
            ws.cell(row=r, column=cols["remark"]).value = remark
        _save(wb, path)
    finally:
        wb.close()


def delete_secret(path: Path, name: str) -> None:
    if not path.exists():
        raise SecretError("not_found", "凭证文件不存在")
    wb, ws, cols = _ensure(path)
    try:
        r = _find_row(ws, cols["name"], (name or "").strip())
        if not r:
            raise SecretError("not_found", f"变量不存在: {name}")
        ws.delete_rows(r)
        _save(wb, path)
    finally:
        wb.close()


def _save(wb, path: Path) -> None:
    """改动写入 ~$ 前缀临时文件，经 load_secret_vars 结构校验后原子替换原文件。"""
    tmp = path.with_name(f"~$platform-write.{path.stem}.tmp.xlsx")
    try:
        wb.save(tmp)
        excel_loader.load_secret_vars(tmp)      # 缺表/缺表头会抛 ExcelFormatError
    except ExcelFormatError as e:
        _safe_remove(tmp)
        raise SecretError("invalid", str(e).replace(tmp.name, path.name))
    except PermissionError:
        _safe_remove(tmp)
        raise SecretError("locked", f"文件 {path.name} 正被其他程序占用（如 Excel 打开），请关闭后重试")
    except Exception as e:  # noqa: BLE001
        _safe_remove(tmp)
        raise SecretError("invalid", f"写入临时文件失败: {e!r}")

    try:
        os.replace(tmp, path)
    except PermissionError:
        _safe_remove(tmp)
        raise SecretError("locked", f"文件 {path.name} 正被其他程序占用（如 Excel 打开），请关闭后重试")


def _safe_remove(p: Path) -> None:
    try:
        p.unlink(missing_ok=True)
    except OSError:
        pass
