"""在 CI（Jenkins）环境中，从环境变量生成 data/secret_vars.xlsx。

框架运行时从 config.yaml 的 secret_vars_file 指向的文件加载凭证，该文件被 .gitignore
忽略、不入仓库。CI 上没有本地文件时，用本脚本把 Jenkins 凭据（以环境变量注入）落盘成
同名 Excel，从而复用完全相同的数据驱动加载链路。

用法（由 Jenkinsfile 调用，凭据通过 withCredentials / 环境变量注入）:
    set API_TOKEN=...        (Linux: export API_TOKEN=...)
    python tools/gen_secret.py

支持的环境变量（缺省时写入占位符，便于本地试跑）:
    API_TOKEN       登录态完整 Authorization 头值（形如 "Basic xxxx"）
    API_USER_ID     业务 userId
    API_USERNAME    账号（可选，覆盖用例全局变量）
    API_PASSWORD    密码（可选，覆盖用例全局变量）
    SECRET_VARS_FILE 输出路径（可选，默认 data/secret_vars.xlsx）
"""
import os
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_OUTPUT = ROOT / "data" / "secret_vars.xlsx"

PLACEHOLDER = "<未从环境变量注入>"


def _rows() -> list:
    """(变量名, 值, 备注)。值来自环境变量，未设置则填占位符。"""
    return [
        ("token", os.environ.get("API_TOKEN", PLACEHOLDER),
         "登录态完整 Authorization 头值（形如 Basic xxx），CI 由凭据注入"),
        ("user_id", os.environ.get("API_USER_ID", PLACEHOLDER), "业务 userId，CI 由凭据注入"),
        ("username", os.environ.get("API_USERNAME", ""), "测试账号（可选，留空不覆盖用例）"),
        ("password", os.environ.get("API_PASSWORD", ""), "测试密码（可选，留空不覆盖用例）"),
    ]


def build(env_rows: list) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "全局变量"
    header_fill = PatternFill("solid", fgColor="C00000")
    header_font = Font(bold=True, color="FFFFFF")
    for col, header in enumerate(["变量名", "值", "备注"], start=1):
        c = ws.cell(row=1, column=col, value=header)
        c.fill = header_fill
        c.font = header_font
    wrap = Alignment(vertical="center", wrap_text=True)
    for row_no, (name, value, note) in enumerate(env_rows, start=2):
        ws.cell(row=row_no, column=1, value=name)
        ws.cell(row=row_no, column=2, value=value).alignment = wrap
        ws.cell(row=row_no, column=3, value=note).alignment = wrap
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 50
    return wb


def main() -> None:
    out = Path(os.environ.get("SECRET_VARS_FILE", DEFAULT_OUTPUT))
    if not out.is_absolute():
        out = ROOT / out
    out.parent.mkdir(parents=True, exist_ok=True)

    env_rows = _rows()
    build(env_rows).save(out)

    # 只打印变量名与是否命中，绝不打印真实值，避免泄露到构建日志
    injected = [n for n, v, _ in env_rows if v and v != PLACEHOLDER]
    print(f"[gen_secret] 已生成凭证文件: {out}")
    print(f"[gen_secret] 从环境变量注入的变量: {injected or '无（全部占位）'}")


if __name__ == "__main__":
    sys.exit(main())
