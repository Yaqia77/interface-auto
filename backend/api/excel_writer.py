"""用例写回：source+row 定位、case_id 行锁、先写临时文件全量校验再原子替换。

安全约定：
- 校验不通过绝不触碰原文件（临时文件方案，excel_loader 是唯一校验口径）；
- 写入前比对行内用例编号与预期一致，防止前端缓存的行号漂移误改他行；
- 目标文件被 Excel 打开锁定时 os.replace 抛 PermissionError → 转 409 提示。
"""
import json
import os
import re
from pathlib import Path

from openpyxl import load_workbook

from core import excel_loader
from core.excel_loader import CASE_COLUMNS, CASE_SHEET, ExcelFormatError

FIELD_TO_HEADER = {f: h for h, f in CASE_COLUMNS.items()}

# 可编辑字段（case_id 不可改：它是行锁定位的锚点；编号变更请走“新增+停用”）
EDITABLE_FIELDS = {"module", "api_name", "case_desc", "priority", "run", "method",
                   "url", "headers", "params", "body", "extract", "asserts", "timeout",
                   "pre_scripts", "post_scripts"}

# 后新增的可选列：存量 Excel 没这两列，首次真正用到时才就地补表头，
# 不让用户为了引用一个脚本而手工重建用例文件。
_AUTO_HEADER_FIELDS = ("pre_scripts", "post_scripts")


class CaseWriteError(Exception):
    """kind: invalid（字段值非法）/ not_found（定位不到行）/ conflict（行锁不匹配）/ locked（文件被占用）"""

    def __init__(self, kind: str, message: str):
        super().__init__(message)
        self.kind = kind


def _serialize(field: str, value):
    """Python 值 → Excel 单元格文本（与存量文件书写习惯一致）。"""
    if field == "run":
        if not isinstance(value, bool):
            raise CaseWriteError("invalid", "run 必须是布尔值")
        return "Y" if value else "N"
    if field == "priority":
        v = str(value or "").strip().upper()
        if v not in excel_loader.VALID_PRIORITIES:
            raise CaseWriteError("invalid", f"优先级必须是 {'/'.join(sorted(excel_loader.VALID_PRIORITIES))}，实际为 {v!r}")
        return v
    if field == "method":
        v = str(value or "").strip().upper()
        if v not in excel_loader.VALID_METHODS:
            raise CaseWriteError("invalid", f"请求方式必须是 {'/'.join(sorted(excel_loader.VALID_METHODS))}")
        return v
    if field == "timeout":
        if value is None or value == "":
            return None
        try:
            f = float(value)
        except (TypeError, ValueError):
            raise CaseWriteError("invalid", f"超时必须是正数（秒），实际为 {value!r}")
        if f <= 0:
            raise CaseWriteError("invalid", "超时必须大于 0")
        return f
    if field in ("headers", "params", "extract"):
        if value in (None, {}):
            return None
        if not isinstance(value, dict):
            raise CaseWriteError("invalid", f"{FIELD_TO_HEADER[field]} 必须是 JSON 对象")
        return json.dumps(value, ensure_ascii=False)
    if field == "body":
        if value in (None, ""):
            return None
        if isinstance(value, str):
            return value
        return json.dumps(value, ensure_ascii=False)
    if field == "asserts":
        if not isinstance(value, list) or not value:
            raise CaseWriteError("invalid", "断言必须是非空数组，如 [[\"eq\",\"$.code\",0]]")
        return json.dumps(value, ensure_ascii=False)
    if field in _AUTO_HEADER_FIELDS:
        # 前置/后置脚本：列表 → 逗号分隔单元格文本（与 excel_loader.split_refs 无损往返）
        # None / 空列表 = 不引用任何脚本（脚手架新建项目带这两列，用例未配脚本时须容忍空，否则整条被拒）
        if value is None:
            return None
        if not isinstance(value, list):
            raise CaseWriteError("invalid",
                                 f"{FIELD_TO_HEADER[field]} 必须是脚本名数组，如 [\"登录取token\"]")
        if not all(isinstance(v, str) for v in value):
            raise CaseWriteError("invalid", f"{FIELD_TO_HEADER[field]} 只能放脚本名字符串")
        names = list(dict.fromkeys(v.strip() for v in value if v.strip()))
        joined = excel_loader.join_refs(names)
        if not joined:
            return None
        if excel_loader.split_refs(joined) != names:
            raise CaseWriteError("invalid", f"脚本名不能含逗号等分隔符: {value}")
        return joined
    # 文本类：module / api_name / case_desc / url
    v = str(value or "").strip()
    if not v:
        raise CaseWriteError("invalid", f"{FIELD_TO_HEADER[field]} 不能为空")
    return v


def apply_update(file_path: Path, case_id: str, row: int, changes: dict) -> None:
    """把 changes（字段名→Python 值）写入 file_path「用例」表第 row 行，校验通过才落盘。"""
    unknown = set(changes) - EDITABLE_FIELDS
    if unknown:
        raise CaseWriteError("invalid", f"不可编辑字段: {', '.join(sorted(unknown))}")
    if not changes:
        raise CaseWriteError("invalid", "没有需要更新的字段")
    if not file_path.is_file():
        raise CaseWriteError("not_found", f"用例文件不存在: {file_path.name}")

    wb = load_workbook(file_path)
    try:
        ws, field_to_col = _sheet_of(wb, file_path)
        _assert_row_lock(ws, field_to_col, case_id, row)

        for field, value in changes.items():
            col = field_to_col.get(field)
            if col is None:
                if field not in _AUTO_HEADER_FIELDS:
                    # 可选列在存量表头中不存在（如「超时(秒)」列缺失）：拒绝而非静默丢改动
                    raise CaseWriteError("invalid",
                                         f"文件 {file_path.name} 没有「{FIELD_TO_HEADER[field]}」列，无法保存该字段")
                if not value:
                    continue        # 该列本就不存在，且本次是清空引用：无事可做
                col = _ensure_column(ws, field_to_col, field)
            ws.cell(row=row, column=col).value = _serialize(field, value)

        _save_validated(wb, file_path)
    finally:
        wb.close()


def append_case(file_path: Path, data: dict) -> None:
    """追加一条新用例到「用例」表末行（编号唯一；同一条校验链路不落盘）。"""
    if not file_path.is_file():
        raise CaseWriteError("not_found", f"用例文件不存在: {file_path.name}")

    case_id = str(data.get("case_id") or "").strip()
    if not case_id:
        raise CaseWriteError("invalid", "用例编号不能为空")
    if re.search(r"[\r\n]", case_id):
        raise CaseWriteError("invalid", "用例编号不能包含换行")
    # 必填口径：编号/接口名/URL/方法/断言；其余字段默认兼容
    for f in ("api_name", "url", "method", "asserts"):
        v = data.get(f)
        if v is None or (isinstance(v, str) and not v.strip()) or (isinstance(v, list) and not v):
            raise CaseWriteError("invalid", f"新增用例缺少必填字段「{FIELD_TO_HEADER[f]}」")
    data.setdefault("module", "未分类")
    data.setdefault("case_desc", "新建用例")
    data.setdefault("priority", "P2")
    data.setdefault("run", True)

    wb = load_workbook(file_path)
    try:
        ws, field_to_col = _sheet_of(wb, file_path)
        # 编号唯一：扫描现列（大小写不敏感，与 loader 取数习惯无关，只防肉眼重复）
        col_id = field_to_col["case_id"]
        existing = {str(ws.cell(row=r, column=col_id).value or "").strip() for r in range(2, ws.max_row + 1)}
        if case_id in existing or case_id.lower() in {e.lower() for e in existing if e}:
            raise CaseWriteError("conflict", f"用例编号已存在: {case_id}（请换一个或编辑现有用例）")

        row = ws.max_row + 1
        ws.cell(row=row, column=col_id).value = case_id
        for field, value in data.items():
            if field == "case_id":
                continue
            if field not in EDITABLE_FIELDS:
                continue  # 未知键忽略（pydantic 已挡，防御性）
            col = field_to_col.get(field)
            if col is None:
                if field in _AUTO_HEADER_FIELDS and value:
                    col = _ensure_column(ws, field_to_col, field)
                else:
                    continue  # 可选列缺失时跳过；必填列缺失在下方 _serialize 后统一报
            ws.cell(row=row, column=col).value = _serialize(field, value)
        for f in ("api_name", "url", "method", "asserts", "module", "case_desc", "priority", "run"):
            if f not in field_to_col:
                raise CaseWriteError("invalid", f"文件 {file_path.name} 缺少「{FIELD_TO_HEADER[f]}」列，无法新增用例")

        _save_validated(wb, file_path)
    finally:
        wb.close()


def delete_case(file_path: Path, case_id: str, row: int) -> None:
    """删除整行（行锁防漂移；删后全量校验通过才落盘）。"""
    if not file_path.is_file():
        raise CaseWriteError("not_found", f"用例文件不存在: {file_path.name}")
    wb = load_workbook(file_path)
    try:
        ws, field_to_col = _sheet_of(wb, file_path)
        _assert_row_lock(ws, field_to_col, case_id, row)
        ws.delete_rows(row)
        _save_validated(wb, file_path)
    finally:
        wb.close()


def _sheet_of(wb, file_path: Path):
    """取「用例」表并解析表头→列号；只校验结构，不改动内容。"""
    if CASE_SHEET not in wb.sheetnames:
        raise CaseWriteError("invalid", f"文件 {file_path.name} 缺少工作表「{CASE_SHEET}」")
    ws = wb[CASE_SHEET]
    field_to_col: dict[str, int] = {}
    for idx, cell in enumerate(ws[1], start=1):
        if cell.value is not None and str(cell.value).strip() in CASE_COLUMNS:
            field_to_col[CASE_COLUMNS[str(cell.value).strip()]] = idx
    if "case_id" not in field_to_col:
        raise CaseWriteError("invalid", "「用例」表缺少「用例编号」列")
    return ws, field_to_col


def _ensure_column(ws, field_to_col: dict[str, int], field: str) -> int:
    """可选列表头缺失时在末列补一个同名列，并登记到 field_to_col。"""
    header = FIELD_TO_HEADER[field]
    col = ws.max_column + 1
    ws.cell(row=1, column=col).value = header
    field_to_col[field] = col
    return col


def _assert_row_lock(ws, field_to_col: dict[str, int], case_id: str, row: int) -> None:
    """行锁：目标行的用例编号必须与预期一致（防行号漂移误改/误删）。"""
    locked_id = str(ws.cell(row=row, column=field_to_col["case_id"]).value or "").strip()
    if locked_id != case_id:
        raise CaseWriteError(
            "conflict",
            f"行定位已漂移：第{row}行当前是用例 {locked_id or '(空)'}，而非 {case_id}。请刷新列表后重试")


def _save_validated(wb, file_path: Path) -> None:
    """内存改动写入 tmp 临时文件，经 excel_loader 全量校验后原子替换原文件。

    tmp 命名：~$ 前缀（loader 目录模式已排除，与 Excel 临时文件同机制）+ .xlsx 后缀（openpyxl 校验扩展名）；
    同目录保证 os.replace 原子性。
    """
    tmp = file_path.with_name(f"~$platform-write.{file_path.stem}.tmp.xlsx")
    try:
        wb.save(tmp)
    except Exception as e:  # noqa: BLE001
        _safe_remove(tmp)
        raise CaseWriteError("invalid", f"写入临时文件失败: {e!r}")

    # 用框架自身的 loader 全量校验临时文件（仅校结构）：allow_empty=True 使“删到只剩表头”也算合法，
    # 避免新建空项目/删完最后一个用例时被“零用例”错误拒绝落盘；通过才原子替换原文件
    try:
        excel_loader.load(tmp, allow_empty=True)
    except ExcelFormatError as e:
        _safe_remove(tmp)
        # 报错定位透出的是临时文件名，映射回原始文件名再给前端
        raise CaseWriteError("invalid", str(e).replace(tmp.name, file_path.name))
    except Exception as e:  # noqa: BLE001（文件损坏等意外一律拒绝落盘）
        _safe_remove(tmp)
        raise CaseWriteError("invalid", f"写回结果校验失败: {e!r}")

    try:
        os.replace(tmp, file_path)          # 同目录重命名，原子操作
    except PermissionError:
        _safe_remove(tmp)
        raise CaseWriteError("locked", f"文件 {file_path.name} 正被其他程序占用（如 Excel 打开），请关闭后重试")


def _safe_remove(p: Path) -> None:
    try:
        p.unlink(missing_ok=True)
    except OSError:
        pass


def bulk_append(file_path: Path, cases: list) -> dict:
    """批量追加用例（一次打开/落盘），已存在的 case_id 跳过。

    cases: TestCase 列表（由 excel_loader.load 解析得到）。
    返回 {"imported": N, "skipped": M, "errors": [str]}。
    """
    if not file_path.is_file():
        raise CaseWriteError("not_found", f"用例文件不存在: {file_path.name}")

    wb = load_workbook(file_path)
    try:
        ws, field_to_col = _sheet_of(wb, file_path)

        # 收集现有 case_id（大小写不敏感）
        col_id = field_to_col["case_id"]
        existing = set()
        for r in range(2, ws.max_row + 1):
            v = str(ws.cell(row=r, column=col_id).value or "").strip()
            if v:
                existing.add(v.lower())

        imported = 0
        skipped = 0
        errors: list[str] = []

        for tc in cases:
            cid = tc.case_id.strip()
            if cid.lower() in existing:
                skipped += 1
                continue
            # 先序列化全部字段，任何失败则整条跳过（不污染工作表）
            row_values: dict[str, object] = {}
            has_error = False
            for field in ("module", "api_name", "case_desc", "priority", "run",
                          "method", "url", "headers", "params", "body",
                          "extract", "asserts", "timeout",
                          "pre_scripts", "post_scripts"):
                val = getattr(tc, field, None)
                col = field_to_col.get(field)
                if col is None:
                    if field in _AUTO_HEADER_FIELDS and val:
                        _ensure_column(ws, field_to_col, field)
                    else:
                        continue
                try:
                    row_values[field] = _serialize(field, val)
                except CaseWriteError as e:
                    errors.append(f"{cid}: {e}")
                    has_error = True
                    break
            if has_error:
                continue

            # 写入新行
            row = ws.max_row + 1
            ws.cell(row=row, column=col_id).value = cid
            existing.add(cid.lower())
            for field, cell_val in row_values.items():
                ws.cell(row=row, column=field_to_col[field]).value = cell_val
            imported += 1

        if imported > 0:
            _save_validated(wb, file_path)
    finally:
        wb.close()

    return {"imported": imported, "skipped": skipped, "errors": errors}
