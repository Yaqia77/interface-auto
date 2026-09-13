"""Excel 用例文件解析与格式校验（报错精确定位到文件、Sheet、行、列）。"""
import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .models import TestCase

CASE_SHEET = "用例"
VAR_SHEET = "全局变量"

# 「用例」表表头 → 字段名
CASE_COLUMNS = {
    "用例编号": "case_id",
    "模块": "module",
    "接口名称": "api_name",
    "用例描述": "case_desc",
    "优先级": "priority",
    "是否执行": "run",
    "请求方式": "method",
    "请求路径": "url",
    "请求头": "headers",
    "查询参数": "params",
    "请求体": "body",
    "变量提取": "extract",
    "断言": "asserts",
    "超时(秒)": "timeout",
}
REQUIRED_HEADERS = ["用例编号", "模块", "接口名称", "用例描述", "优先级", "是否执行", "请求方式", "请求路径", "断言"]
VALID_METHODS = {"GET", "POST", "PUT", "DELETE", "PATCH"}
VALID_PRIORITIES = {"P0", "P1", "P2", "P3"}
ASSERT_OPS = {"eq", "neq", "contains", "gt", "ge", "lt", "le", "regex", "not_null", "len_eq"}
OPS_NEED_EXPECTED = ASSERT_OPS - {"not_null"}
RUN_TRUE = {"Y", "YES", "TRUE", "是"}
RUN_FALSE = {"N", "NO", "FALSE", "否"}


class ExcelFormatError(Exception):
    """Excel 内容/格式不合法，message 中包含全部错误的精确位置。"""


class LoadResult:
    """一次加载的结果：用例列表 + 全局变量。"""

    def __init__(self, cases: list, variables: dict):
        self.cases = cases
        self.variables = variables


def load(path) -> LoadResult:
    """加载用例文件；path 可为单个 .xlsx 文件或目录（目录下所有 .xlsx 合并，按文件名排序）。"""
    p = Path(path)
    if p.is_dir():
        files = sorted(p.glob("*.xlsx"))
        files = [f for f in files if not f.name.startswith("~$")]  # 排除 Excel 打开时的临时文件
        if not files:
            raise FileNotFoundError(f"目录 {p} 下没有 .xlsx 用例文件")
    else:
        if not p.exists():
            raise FileNotFoundError(f"用例文件不存在: {p}")
        files = [p]

    cases: list[TestCase] = []
    variables: dict = {}
    for f in files:
        _load_one(f, cases, variables)
    if not cases:
        raise ExcelFormatError(f"未从 {len(files)} 个文件中解析到任何用例")
    return LoadResult(cases, variables)


def _load_one(path: Path, cases: list, variables: dict) -> None:
    errors: list[str] = []
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        if CASE_SHEET not in wb.sheetnames:
            errors.append(f"[{path.name}] 缺少工作表「{CASE_SHEET}」")
            raise ExcelFormatError(_join(errors))
        ws = wb[CASE_SHEET]

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise ExcelFormatError(_join([f"[{path.name} → {CASE_SHEET}] 表格为空"]))

        # 表头定位（列顺序不限）
        header_to_field = {}
        field_to_col = {}  # 字段 -> 0 基列号
        for idx, cell in enumerate(rows[0]):
            if cell is None:
                continue
            header = str(cell).strip()
            if header in CASE_COLUMNS:
                if header in header_to_field:
                    errors.append(f"[{path.name} → {CASE_SHEET}] 表头「{header}」重复")
                header_to_field[header] = CASE_COLUMNS[header]
                field_to_col[CASE_COLUMNS[header]] = idx
        missing = [h for h in REQUIRED_HEADERS if h not in header_to_field]
        if missing:
            errors.append(f"[{path.name} → {CASE_SHEET}] 缺少必填表头: {', '.join(missing)}")
            raise ExcelFormatError(_join(errors))

        # 逐行解析用例
        seen_ids: dict[str, str] = {}
        for row_no, row in enumerate(rows[1:], start=2):
            raw = {}
            for field, col in field_to_col.items():
                raw[field] = row[col] if col < len(row) else None
            if all(v is None or (isinstance(v, str) and not v.strip()) for v in raw.values()):
                continue  # 整行为空，跳过
            case = _parse_case(path.name, row_no, raw, field_to_col, seen_ids, errors)
            if case is not None:
                cases.append(case)

        # 全局变量表（可选）
        if VAR_SHEET in wb.sheetnames:
            _load_variables(path.name, wb[VAR_SHEET], variables, errors)
    finally:
        wb.close()

    if errors:
        raise ExcelFormatError(_join(errors))


def _parse_case(file_name, row_no, raw, field_to_col, seen_ids, errors):
    def cell_text(field):
        v = raw.get(field)
        if v is None:
            return ""
        return v.strip() if isinstance(v, str) else v

    def err(field, msg):
        col = field_to_col.get(field)
        header = next((h for h, f in CASE_COLUMNS.items() if f == field), field)
        pos = f"第{row_no}行 {get_column_letter(col + 1)}列({header})" if col is not None else f"第{row_no}行"
        errors.append(f"[{file_name} → {CASE_SHEET} {pos}] {msg}")

    case_id = str(cell_text("case_id") or "").strip()
    if not case_id:
        err("case_id", "用例编号不能为空")
        return None
    if case_id in seen_ids:
        err("case_id", f"用例编号 {case_id} 重复（首次出现: {seen_ids[case_id]}）")
        return None
    seen_ids[case_id] = f"第{row_no}行"

    for field in ("module", "api_name", "case_desc", "url"):
        if not str(cell_text(field) or "").strip():
            err(field, f"{dict((v, k) for k, v in CASE_COLUMNS.items())[field]}不能为空")
            return None

    priority = str(cell_text("priority") or "").strip().upper()
    if priority not in VALID_PRIORITIES:
        err("priority", f"优先级必须是 {'/'.join(sorted(VALID_PRIORITIES))}，实际为 {priority!r}")
        return None

    run_raw = str(cell_text("run") or "").strip().upper()
    if run_raw in RUN_TRUE:
        run = True
    elif run_raw in RUN_FALSE:
        run = False
    else:
        err("run", f"是否执行必须是 Y/N，实际为 {run_raw!r}")
        return None

    method = str(cell_text("method") or "").strip().upper()
    if method not in VALID_METHODS:
        err("method", f"请求方式必须是 {'/'.join(sorted(VALID_METHODS))}，实际为 {method!r}")
        return None

    headers = _parse_json_object(cell_text("headers"), "headers", err)
    params = _parse_json_object(cell_text("params"), "params", err)
    extract = _parse_json_object(cell_text("extract"), "extract", err)
    body = _parse_body(cell_text("body"), err)
    asserts = _parse_asserts(cell_text("asserts"), err)

    timeout_raw = cell_text("timeout")
    timeout = None
    if timeout_raw not in (None, ""):
        try:
            timeout = float(timeout_raw)
            if timeout <= 0:
                raise ValueError
        except (TypeError, ValueError):
            err("timeout", f"超时必须是正数（秒），实际为 {timeout_raw!r}")
            return None

    ok_so_far = headers is not False and params is not False and extract is not False and asserts is not False and body is not False
    if not ok_so_far:
        return None  # 该行已有 JSON 相关错误，跳过

    return TestCase(
        row=row_no,
        source=file_name,
        case_id=case_id,
        module=str(raw["module"]).strip(),
        api_name=str(raw["api_name"]).strip(),
        case_desc=str(raw["case_desc"]).strip(),
        priority=priority,
        run=run,
        method=method,
        url=str(raw["url"]).strip(),
        headers=headers,
        params=params,
        body=body,
        extract=extract,
        asserts=asserts,
        timeout=timeout,
    )


def _parse_json_object(value, field, err):
    """解析必须为 JSON 对象的列，返回 dict / None；格式错误时记录 err 并返回 False 哨兵。"""
    if value in (None, ""):
        return None
    try:
        parsed = json.loads(value) if isinstance(value, str) else value
    except json.JSONDecodeError as e:
        err(field, f"JSON 解析失败: {e}")
        return False
    if not isinstance(parsed, dict):
        err(field, f"必须是 JSON 对象 {{...}}，实际是 {type(parsed).__name__}")
        return False
    return parsed


def _parse_body(value, err):
    """请求体：以 { 或 [ 开头时必须为合法 JSON；否则按原始文本处理。"""
    if value in (None, ""):
        return None
    text = value if isinstance(value, str) else str(value)
    stripped = text.strip()
    if stripped[:1] in ("{", "["):
        try:
            return json.loads(stripped)
        except json.JSONDecodeError as e:
            err("body", f"JSON 解析失败: {e}")
            return False
    return stripped  # 非 JSON 文本，按原始字符串请求体处理


def _parse_asserts(value, err):
    """断言列：JSON 数组，元素为 [操作符, 取值路径, 期望值]。"""
    if value in (None, ""):
        err("asserts", "断言不能为空")
        return False
    try:
        rules = json.loads(value) if isinstance(value, str) else value
    except json.JSONDecodeError as e:
        err("asserts", f"JSON 解析失败: {e}")
        return False
    if not isinstance(rules, list) or not rules:
        err("asserts", "必须是 JSON 数组且至少一条，如 [[\"eq\",\"status_code\",200]]")
        return False
    for i, rule in enumerate(rules, 1):
        if not isinstance(rule, (list, tuple)) or not 2 <= len(rule) <= 3:
            err("asserts", f"第{i}条格式错误，应为 [操作符, 取值路径, 期望值]，如 [\"eq\",\"$.code\",0]")
            return False
        op = str(rule[0]).strip().lower()
        if op not in ASSERT_OPS:
            err("asserts", f"第{i}条操作符 {rule[0]!r} 不支持，可用: {', '.join(sorted(ASSERT_OPS))}")
            return False
        if len(rule) < 3 and op in OPS_NEED_EXPECTED:
            err("asserts", f"第{i}条操作符 {op} 需要期望值")
            return False
        if not str(rule[1] or "").strip():
            err("asserts", f"第{i}条取值路径不能为空")
            return False
    return rules


def _load_variables(file_name, ws, variables: dict, errors: list) -> None:
    """解析「全局变量」表：变量名 / 值 / 备注。"""
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return
    header_idx = {}
    for idx, cell in enumerate(rows[0]):
        if cell is not None and str(cell).strip() in ("变量名", "值"):
            header_idx[str(cell).strip()] = idx
    if "变量名" not in header_idx or "值" not in header_idx:
        errors.append(f"[{file_name} → {VAR_SHEET} 表头] 需包含「变量名」和「值」两列")
        return
    for row_no, row in enumerate(rows[1:], start=2):
        name = row[header_idx["变量名"]] if header_idx["变量名"] < len(row) else None
        value = row[header_idx["值"]] if header_idx["值"] < len(row) else None
        name = str(name).strip() if name is not None else ""
        if not name and value in (None, ""):
            continue
        if not name:
            errors.append(f"[{file_name} → {VAR_SHEET} 第{row_no}行] 变量名为空")
            continue
        if name in variables:
            errors.append(f"[{file_name} → {VAR_SHEET} 第{row_no}行] 变量 {name} 重复定义")
            continue
        if value is None:
            value = ""
        variables[name] = value if isinstance(value, str) else value


def _join(errors: list) -> str:
    return "Excel 用例文件校验失败，共 " + str(len(errors)) + " 处错误:\n" + "\n".join(errors)
