# -*- coding: utf-8 -*-
"""UI 用例 Excel 一步一行「下载模板 → 填例 → 预览 → 导入」端到端冒烟。

覆盖低门槛录入方案一（Excel 批量导入）与导出回导闭环的接口链路，不依赖浏览器（纯数据链路）：
1. /template?format=excel 下载 .xlsx 模板（zip 魔数 PK）；
2. 程序化在模板里填一条真实用例（中文动作：打开/输入/点击/断言文本），删掉示例行；
3. /preview 上传解析：一步一行按编号分组为 1 用例、动作归一为英文键、importable=true；
4. /import 提交预览产出的 cases：落库后可从 /api/ui-cases 列表回读，动作已规范；
5. 重复导入同编号 → skipped（查重）；示例行编号前缀不污染真实数据；
6. /export?format=excel|json 导出：Excel 与模板同构可回导（exists 查重生效），JSON 与 /import 同构无损。

用法（后端需带 playwright 的 venv 启动，本冒烟只用数据链路不驱动浏览器）：
    .venv\Scripts\python.exe -m uvicorn backend.main:app --port 8000
    .venv\Scripts\python.exe tools\smoke_ui_excel.py
只在临时 UI 项目造数据，收尾必删。
"""
import io
import json
import os
import sys
import urllib.error
import urllib.parse
import urllib.request
import uuid
from pathlib import Path

from openpyxl import load_workbook

BASE = os.environ.get("IA_BASE", "http://127.0.0.1:8000").rstrip("/")
ADMIN = ("admin", "admin123")

_passed = _failed = 0


def check(name, cond, extra=""):
    global _passed, _failed
    if cond:
        _passed += 1
        print(f"  [OK] {name}")
    else:
        _failed += 1
        print(f"  [FAIL] {name} {extra}")


def req(method, path, body=None, token=None):
    data = json.dumps(body).encode() if body is not None else None
    r = urllib.request.Request(BASE + path, data=data, method=method)
    r.add_header("Content-Type", "application/json")
    if token:
        r.add_header("Authorization", f"Bearer {token}")
    try:
        with urllib.request.urlopen(r) as resp:
            return resp.status, json.loads(resp.read().decode() or "{}")
    except urllib.error.HTTPError as e:
        raw = e.read()
        try:
            return e.code, json.loads(raw.decode() or "{}")
        except json.JSONDecodeError:
            return e.code, {"raw": raw.decode()[:200]}
    except urllib.error.URLError as e:
        print(f"后端不可达：{BASE}（{e.reason}）", file=sys.stderr)
        sys.exit(2)


def get_bytes(path, token=None):
    r = urllib.request.Request(BASE + path, method="GET")
    if token:
        r.add_header("Authorization", f"Bearer {token}")
    with urllib.request.urlopen(r) as resp:
        return resp.status, resp.read()


def post_multipart(path, field, filename, content, token=None):
    """手工构造 multipart/form-data 上传（预览 Excel）。"""
    boundary = "----IAExcelSmoke" + uuid.uuid4().hex
    body = io.BytesIO()
    body.write(f"--{boundary}\r\n".encode())
    body.write(f'Content-Disposition: form-data; name="{field}"; filename="{filename}"\r\n'
               .encode())
    body.write(b"Content-Type: application/octet-stream\r\n\r\n")
    body.write(content)
    body.write(f"\r\n--{boundary}--\r\n".encode())
    data = body.getvalue()
    r = urllib.request.Request(BASE + path, data=data, method="POST")
    r.add_header("Content-Type", f"multipart/form-data; boundary={boundary}")
    if token:
        r.add_header("Authorization", f"Bearer {token}")
    try:
        with urllib.request.urlopen(r) as resp:
            return resp.status, json.loads(resp.read().decode() or "{}")
    except urllib.error.HTTPError as e:
        raw = e.read()
        try:
            return e.code, json.loads(raw.decode() or "{}")
        except json.JSONDecodeError:
            return e.code, {"raw": raw.decode()[:200]}


def login(u, p):
    code, resp = req("POST", "/api/auth/login", {"username": u, "password": p})
    assert code == 200, f"登录失败 {u}: {code} {resp}"
    return resp["token"]


def qf(**extra):
    return "?" + urllib.parse.urlencode({k: v for k, v in extra.items() if v is not None})


def build_filled_xlsx(token) -> bytes:
    """下载模板 → 删除示例数据行 → 追加一条真实用例（中文动作、一步一行）。"""
    code, raw = get_bytes("/api/ui-cases/template?format=excel", token)
    assert code == 200, f"下载 Excel 模板失败 {code}"
    wb = load_workbook(io.BytesIO(raw))
    ws = wb["UI步骤"]
    # 清掉除表头外的所有示例行
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    # 列：用例编号 名称 模块 优先级 浏览器 无头 是否执行 重试 步骤号 动作 定位器 取值 备注
    uniq = BUILD_TAG
    ws.append([f"wx_{uniq}", "冒烟登录", "冒烟", "P1", "chromium", "是", "Y", "0",
               "1", "打开", "", "/login", "打开登录页"])
    ws.append(["", "", "", "", "", "", "", "", "2", "输入", "#username", "qa", ""])
    ws.append(["", "", "", "", "", "", "", "", "3", "点击", "#submit", "", ""])
    ws.append(["", "", "", "", "", "", "", "", "4", "断言文本", ".welcome", "欢迎", "校验成功"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


BUILD_TAG = ""


def main():
    global BUILD_TAG
    print(f"目标：{BASE}")
    code, _ = get_bytes("/api/health")
    check("健康检查可访问", code == 200, code)
    tok = login(*ADMIN)

    # 1) 下载模板
    code, raw = get_bytes("/api/ui-cases/template?format=excel", tok)
    check("下载 Excel 模板 200 + xlsx 魔数(PK)", code == 200 and raw[:2] == b"PK", code)

    # 2) 建临时 UI 项目
    BUILD_TAG = uuid.uuid4().hex[:6]
    code, proj = req("POST", "/api/projects", {"name": f"Excel冒烟_{BUILD_TAG}", "kind": "ui"}, tok)
    check("建 UI 项目 201", code == 201 and proj.get("file", "").endswith(".json"), f"{code} {proj}")
    pfile = proj.get("file")
    pid = proj.get("id", "")
    if not pfile:
        return 1

    try:
        content = build_filled_xlsx(tok)

        # 3) 上传预览
        code, prev = post_multipart("/api/ui-cases/preview" + qf(project=pfile),
                                    "file", "cases.xlsx", content, tok)
        check("预览解析 200", code == 200, f"{code} {prev}")
        items = prev.get("items") or []
        check("解析出 1 条用例（一步一行按编号分组）", len(items) == 1, prev.get("total"))
        if items:
            it = items[0]
            check("步骤数=4", it.get("steps_count") == 4, it.get("steps_count"))
            check("该条可导入 importable=true", it.get("importable") is True, it.get("marks"))
            check("未与已有用例重复 exists=false", it.get("exists") is False, it.get("exists"))
            acts = [s.get("action") for s in (it.get("case", {}).get("steps") or [])]
            check("中文动作已归一为英文键",
                  acts == ["goto", "fill", "click", "expect_text"], acts)

        # 4) 提交导入（复用 /import，cases 来自预览）
        cases = [it["case"] for it in items if it.get("importable")]
        code, imp = req("POST", "/api/ui-cases/import" + qf(file=pfile), {"cases": cases}, tok)
        check("导入 1 条成功", code == 200 and imp.get("imported") == 1, f"{code} {imp}")

        # 5) 列表回读，动作规范、字段完整
        want = f"wx_{BUILD_TAG}"
        code, listing = req("GET", "/api/ui-cases" + qf(file=pfile), token=tok)
        got = {c.get("case_id"): c for c in (listing.get("items") or [])}
        check("列表回读命中该用例", want in got, list(got))
        if want in got:
            c = got[want]
            check("回读用例名/模块正确", c.get("name") == "冒烟登录" and c.get("module") == "冒烟",
                  f"{c.get('name')}/{c.get('module')}")
            check("回读步骤动作为规范英文键",
                  [s.get("action") for s in (c.get("steps") or [])] == ["goto", "fill", "click", "expect_text"],
                  c.get("steps"))

        # 6) 重复导入同编号 → skipped（查重不阻断）
        code, imp2 = req("POST", "/api/ui-cases/import" + qf(file=pfile), {"cases": cases}, tok)
        check("重复导入被跳过（skipped）", code == 200 and want in (imp2.get("skipped") or [])
              and imp2.get("imported") == 0, f"{code} {imp2}")

        # 6.1) 导出 Excel：同模板结构可回导，回导同项目时 exists 查重生效
        ecode, eraw = get_bytes("/api/ui-cases/export" + qf(file=pfile, format="excel"), tok)
        check("导出 Excel 200 + xlsx 魔数(PK)", ecode == 200 and eraw[:2] == b"PK", ecode)
        pcode2, prev2 = post_multipart("/api/ui-cases/preview" + qf(project=pfile),
                                       "file", "export.xlsx", eraw, tok)
        eitems = prev2.get("items") or []
        check("导出件再预览：命中刚才导入的用例且 exists=true（回导会跳过）",
              pcode2 == 200 and any(i.get("case_id") == want and i.get("exists") for i in eitems),
              f"{pcode2} {[(i.get('case_id'), i.get('exists')) for i in eitems]}")

        # 6.2) 导出 JSON：{cases:[...]} 与导入同构，含目标用例且步骤完整（无损）
        jcode, jraw = get_bytes("/api/ui-cases/export" + qf(file=pfile, format="json"), tok)
        jok = jcode == 200
        jc = {}
        if jok:
            try:
                jc = next((c for c in json.loads(jraw).get("cases") or []
                           if c.get("case_id") == want), {})
            except json.JSONDecodeError:
                jok = False
        check("导出 JSON 200 且含该用例（同构可回导）",
              jok and bool(jc), jcode)
        check("导出 JSON 步骤动作已规范、条数一致",
              [s.get("action") for s in (jc.get("steps") or [])]
              == ["goto", "fill", "click", "expect_text"], jc.get("steps"))

        # 7) 预览对空/非法文件的容错（示例-only 模板应解析 0 条）
        tpl_code, tpl_raw = get_bytes("/api/ui-cases/template?format=excel", tok)
        pcode, prev0 = post_multipart("/api/ui-cases/preview" + qf(project=pfile),
                                      "file", "tpl.xlsx", tpl_raw, tok)
        check("纯示例模板预览解析 0 条（示例行被跳过）",
              pcode == 200 and (prev0.get("total") == 0), f"{pcode} total={prev0.get('total')}")
    finally:
        code, resp = req("DELETE", "/api/projects/" + urllib.parse.quote(pid, safe=""), token=tok)
        check("删除临时 UI 项目", code == 200, f"{code} {resp}")

    print(f"\n结果：通过 {_passed} / 失败 {_failed}")
    return 1 if _failed else 0


if __name__ == "__main__":
    sys.exit(main())
