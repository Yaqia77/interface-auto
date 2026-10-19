"""用例写回：source+row 定位、case_id 行锁、先写临时文件全量校验再原子替换。

安全约定：
- 校验不通过绝不触碰原文件（临时文件方案，excel_loader 是唯一校验口径）；
- 写入前比对行内用例编号与预期一致，防止前端缓存的行号漂移误改他行；
- 目标文件被 Excel 打开锁定时 os.replace 抛 PermissionError → 转 409 提示。
"""
import json
import os
from pathlib import Path

from openpyxl import load_workbook

from core import excel_loader
from core.excel_loader import CASE_COLUMNS, CASE_SHEET, ExcelFormatError

FIELD_TO_HEADER = {f: h for h, f in CASE_COLUMNS.items()}

# 可编辑字段（case_id 不可改：它是行锁定位的锚点；编号变更请走“新增+停用”）
EDITABLE_FIELDS = {"module", "api_name", "case_desc", "priority", "run", "method",
                   "url", "headers", "params", "body", "extract", "asserts", "timeout"}


class CaseWriteError(Exception):
    """kind: invalid（字段值非法）/ not_found（定位不到行）/ conflict（行锁不匹配）/ locked（文件被占用）"""

    def __init__(self, kind: str, message: str):
        super().__init__(message)
        self.kind = kind


def _serialize(field: str, value):
    """Python 值 → Excel 单元格文本（与存量文件书写习惯一致）。"""
    if field == "run":
        if not isinstance(value, bool):
            raise CaseWriteError("invalid", "run 必须是布尔值")
        return "Y" if value else "N"
    if field == "priority":
        v = str(value or "").strip().upper()
        if v not in excel_loader.VALID_PRIORITIES:
            raise CaseWriteError("invalid", f"优先级必须是 {'/'.join(sorted(excel_loader.VALID_PRIORITIES))}，实际为 {v!r}")
        return v
    if field == "method":
        v = str(value or "").strip().upper()
        if v not in excel_loader.VALID_METHODS:
            raise CaseWriteError("invalid", f"请求方式必须是 {'/'.join(sorted(excel_loader.VALID_METHODS))}")
        return v
    if field == "timeout":
        if value is None or value == "":
            return None
        try:
            f = float(value)
        except (TypeError, ValueError):
            raise CaseWriteError("invalid", f"超时必须是正数（秒），实际为 {value!r}")
        if f <= 0:
            raise CaseWriteError("invalid", "超时必须大于 0")
        return f
    if field in ("headers", "params", "extract"):
        if value in (None, {}):
            return None
        if not isinstance(value, dict):
            raise CaseWriteError("invalid", f"{FIELD_TO_HEADER[field]} 必须是 JSON 对象")
        return json.dumps(value, ensure_ascii=False)
    if field == "body":
        if value in (None, ""):
            return None
        if isinstance(value, str):
            return value
        return json.dumps(value, ensure_ascii=False)
    if field == "asserts":
        if not isinstance(value, list) or not value:
            raise CaseWriteError("invalid", "断言必须是非空数组，如 [[\"eq\",\"$.code\",0]]")
        return json.dumps(value, ensure_ascii=False)
    # 文本类：module / api_name / case_desc / url
    v = str(value or "").strip()
    if not v:
        raise CaseWriteError("invalid", f"{FIELD_TO_HEADER[field]} 不能为空")
    return v


def apply_update(file_path: Path, case_id: str, row: int, changes: dict) -> None:
    """把 changes（字段名→Python 值）写入 file_path「用例」表第 row 行，校验通过才落盘。"""
    unknown = set(changes) - EDITABLE_FIELDS
    if unknown:
        raise CaseWriteError("invalid", f"不可编辑字段: {', '.join(sorted(unknown))}")
    if not changes:
        raise CaseWriteError("invalid", "没有需要更新的字段")
    if not file_path.is_file():
        raise CaseWriteError("not_found", f"用例文件不存在: {file_path.name}")

    wb = load_workbook(file_path)
    # tmp 命名：~$ 前缀（loader 目录模式已排除，与 Excel 临时文件同机制）+ .xlsx 后缀（openpyxl 校验扩展名）；
    # 同目录保证 os.replace 原子性
    tmp = file_path.with_name(f"~$platform-write.{file_path.stem}.tmp.xlsx")
    try:
        if CASE_SHEET not in wb.sheetnames:
            raise CaseWriteError("invalid", f"文件 {file_path.name} 缺少工作表「{CASE_SHEET}」")
        ws = wb[CASE_SHEET]

        field_to_col: dict[str, int] = {}
        for idx, cell in enumerate(ws[1], start=1):
            if cell.value is not None and str(cell.value).strip() in CASE_COLUMNS:
                field_to_col[CASE_COLUMNS[str(cell.value).strip()]] = idx
        if "case_id" not in field_to_col:
            raise CaseWriteError("invalid", "「用例」表缺少「用例编号」列")

        # 行锁：目标行的用例编号必须与预期一致（防行号漂移误改）
        locked_id = str(ws.cell(row=row, column=field_to_col["case_id"]).value or "").strip()
        if locked_id != case_id:
            raise CaseWriteError(
                "conflict",
                f"行定位已漂移：第{row}行当前是用例 {locked_id or '(空)'}，而非 {case_id}。请刷新列表后重试")

        for field, value in changes.items():
            col = field_to_col.get(field)
            if col is None:
                # 可选列在存量表头中不存在（如「超时(秒)」列缺失）：拒绝而非静默丢改动
                raise CaseWriteError("invalid", f"文件 {file_path.name} 没有「{FIELD_TO_HEADER[field]}」列，无法保存该字段")
            ws.cell(row=row, column=col).value = _serialize(field, value)

        wb.save(tmp)
    finally:
        wb.close()

    # 用框架自身的 loader 全量校验临时文件，通过才原子替换原文件
    try:
        excel_loader.load(tmp)
    except ExcelFormatError as e:
        _safe_remove(tmp)
        # 报错定位透出的是临时文件名，映射回原始文件名再给前端
        raise CaseWriteError("invalid", str(e).replace(tmp.name, file_path.name))
    except Exception as e:  # noqa: BLE001（文件损坏等意外一律拒绝落盘）
        _safe_remove(tmp)
        raise CaseWriteError("invalid", f"写回结果校验失败: {e!r}")

    try:
        os.replace(tmp, file_path)          # 同目录重命名，原子操作
    except PermissionError:
        _safe_remove(tmp)
        raise CaseWriteError("locked", f"文件 {file_path.name} 正被其他程序占用（如 Excel 打开），请关闭后重试")


def _safe_remove(p: Path) -> None:
    try:
        p.unlink(missing_ok=True)
    except OSError:
        pass
